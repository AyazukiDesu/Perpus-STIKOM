"""Ekspor data ke file Excel (.xlsx) dengan header tebal & lebar kolom otomatis."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def buat_excel(judul_sheet, header, baris):
    """header: list nama kolom | baris: list of list. Kembalikan BytesIO siap dikirim."""
    wb = Workbook()
    ws = wb.active
    ws.title = judul_sheet[:31]
    ws.append(header)
    for sel in ws[1]:
        sel.font = Font(bold=True, color="FFFFFF")
        sel.fill = PatternFill("solid", fgColor="5C211D")
        sel.alignment = Alignment(vertical="center")
    for r in baris:
        ws.append(r)
    for i, kolom in enumerate(ws.columns, start=1):
        terpanjang = max((len(str(c.value)) for c in kolom if c.value is not None), default=8)
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(terpanjang + 3, 50)
    ws.freeze_panes = "A2"
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
